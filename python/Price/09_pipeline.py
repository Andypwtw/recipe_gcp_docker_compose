"""依序執行 MOA 價格資料處理流程。

預設以此檔案所在資料夾為工作資料夾；若腳本放在其他位置，可指定
``--directory /path/to/project``。流程依序執行 01–06 的 MOA 原始資料、
01–06 的篩選、07 的換算與 08 的營養資料比對。所有子腳本都會以同一個
工作資料夾執行，以保留它們原本使用的相對路徑。

範例：
    python pipeline.py --directory /path/to/Price

本程式每次啟動只執行一次；排程由 Airflow DAG 管理。
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


TAIWAN_TIMEZONE = "Asia/Taipei"


def configure_taiwan_timezone() -> None:
    """將本流程與啟動的子程式固定使用台灣時區。"""
    os.environ["TZ"] = TAIWAN_TIMEZONE
    # Linux 的 Airflow 容器支援 tzset；保留 hasattr 以維持跨平台相容性。
    if hasattr(time, "tzset"):
        time.tzset()


@dataclass(frozen=True)
class Stage:
    """一個要執行的腳本群組，以及完成後應出現的 JSON 條件。"""

    name: str
    script_pattern: str
    output_selector: Callable[[Path], bool]
    minimum_outputs: Callable[[int], int]
    exact_moa_outputs: bool = False


def json_snapshot(directory: Path) -> dict[Path, int]:
    """回傳所有 JSON 的最後修改時間，用於辨識本階段新產出或更新的檔案。"""
    return {
        path: path.stat().st_mtime_ns
        for path in directory.glob("*.json")
        if path.is_file()
    }


def changed_outputs(
    directory: Path,
    before: dict[Path, int],
    selector: Callable[[Path], bool],
) -> list[Path]:
    """找出本階段建立或更新、且符合命名規則的 JSON 檔。"""
    after = json_snapshot(directory)
    return sorted(
        path
        for path, modified_at in after.items()
        if selector(path) and (path not in before or before[path] != modified_at)
    )


def run_script(script: Path, directory: Path, dry_run: bool) -> None:
    command = [sys.executable, str(script)]
    print("$", " ".join(command), flush=True)
    if not dry_run:
        subprocess.run(command, cwd=directory, check=True)


def run_stage(stage: Stage, directory: Path, dry_run: bool) -> list[Path]:
    scripts = sorted(path for path in directory.glob(stage.script_pattern) if path.is_file())
    if not scripts:
        raise FileNotFoundError(f"找不到 {stage.script_pattern} 腳本（階段：{stage.name}）")

    print(f"\n=== {stage.name}：{len(scripts)} 個腳本 ===", flush=True)
    before = json_snapshot(directory)
    for script in scripts:
        run_script(script, directory, dry_run)

    if dry_run:
        return []

    if stage.exact_moa_outputs:
        required = [script.with_suffix(".json") for script in scripts]
        missing = [path.name for path in required if not path.is_file()]
        if missing:
            raise RuntimeError(
                "MOA 階段完成但缺少對應 JSON：" + ", ".join(missing)
            )
        outputs = changed_outputs(directory, before, stage.output_selector)
        # 即使來源資料沒有變化，MOA 腳本仍需至少保有每支腳本的同名 JSON。
        print("產出：", ", ".join(path.name for path in required), flush=True)
        return outputs

    outputs = changed_outputs(directory, before, stage.output_selector)
    required_count = stage.minimum_outputs(len(scripts))
    if len(outputs) < required_count:
        raise RuntimeError(
            f"{stage.name} 階段預期至少產生或更新 {required_count} 個符合條件的 JSON，"
            f"實際為 {len(outputs)} 個。"
        )
    print("產出：", ", ".join(path.name for path in outputs), flush=True)
    return outputs


STAGES = (
    Stage(
        name="MOA 原始資料",
        script_pattern="0[1-6]_moa_*.py",
        output_selector=lambda path: "_moa_" in path.stem,
        minimum_outputs=lambda script_count: script_count,
        exact_moa_outputs=True,
    ),
    Stage(
        name="平均價格篩選",
        script_pattern="0[1-6]_filter_*.py",
        output_selector=lambda path: "avgprice" in path.stem and "avgprice_per_100g" not in path.stem,
        minimum_outputs=lambda script_count: script_count,
    ),
    Stage(
        name="每 100g 價格換算",
        script_pattern="07_convert_*.py",
        output_selector=lambda path: "avgprice_per_100g" in path.stem,
        minimum_outputs=lambda _script_count: 1,
    ),
    Stage(
        name="營養資料比對",
        script_pattern="08_match_*.py",
        output_selector=lambda path: path.name.startswith("08_match_") and path.suffix == ".json",
        minimum_outputs=lambda _script_count: 1,
    ),
)


MATCHED_PRICES_FILE = "08_match_nutrition_foods__moa_prices.json"
# Docker Compose 會把 recipe_gcp_docker_compose/data 掛載到容器的 /workspace/data。
CONTAINER_REFERENCE_WORKBOOK = Path("/workspace/data/reference/food_nutrition_2025.xlsx")
LOCAL_REFERENCE_WORKBOOK_RELATIVE_PATH = Path(
    "recipe_gcp_docker_compose/data/reference/food_nutrition_2025.xlsx"
)
REFERENCE_SHEET = "台灣食品成分表"
HEADER_ROW = 2
MATCH_HEADERS = ("樣品名稱", "俗名")
PRICE_HEADER = "每100g的價格"


def resolve_reference_workbook(directory: Path) -> Path:
    """取得容器或本機環境中的 food_nutrition_2025.xlsx 位置。"""
    configured_path = os.environ.get("FOOD_NUTRITION_XLSX")
    candidates = [CONTAINER_REFERENCE_WORKBOOK]
    if configured_path:
        candidates.insert(0, Path(configured_path).expanduser())
    candidates.append(directory.parent / LOCAL_REFERENCE_WORKBOOK_RELATIVE_PATH)

    for candidate in candidates:
        if candidate.is_file():
            return candidate
    searched_paths = "、".join(str(candidate) for candidate in candidates)
    raise FileNotFoundError(f"找不到食材營養資料庫，已查詢：{searched_paths}")


def update_reference_workbook(directory: Path) -> None:
    """將 08_match 的新價格寫入食材營養資料庫，未配對資料保留既有價格。"""
    matched_prices_path = directory / MATCHED_PRICES_FILE
    workbook_path = resolve_reference_workbook(directory)
    if not matched_prices_path.is_file():
        raise FileNotFoundError(f"找不到營養資料比對結果：{matched_prices_path}")

    with matched_prices_path.open("r", encoding="utf-8") as input_file:
        matched_groups = json.load(input_file)

    if not isinstance(matched_groups, dict):
        raise RuntimeError("營養資料比對結果必須是依類別分組的 JSON 物件。")

    prices_by_food: dict[tuple[str, str], float | int] = {}
    for records in matched_groups.values():
        if not isinstance(records, list):
            continue
        for record in records:
            if not isinstance(record, dict):
                continue
            sample_name = str(record.get("樣品名稱", "")).strip()
            common_name = str(record.get("俗名", "")).strip()
            price = record.get(PRICE_HEADER)
            if sample_name and isinstance(price, (int, float)):
                prices_by_food[(sample_name, common_name)] = price

    workbook = load_workbook(workbook_path)
    if REFERENCE_SHEET not in workbook.sheetnames:
        raise RuntimeError(f"活頁簿找不到工作表：{REFERENCE_SHEET}")
    worksheet = workbook[REFERENCE_SHEET]
    headers = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[HEADER_ROW]
        if cell.value is not None
    }
    missing_headers = [header for header in (*MATCH_HEADERS, PRICE_HEADER) if header not in headers]
    if missing_headers:
        raise RuntimeError("活頁簿缺少欄位：" + "、".join(missing_headers))

    sample_column, common_name_column = (headers[header] for header in MATCH_HEADERS)
    price_column = headers[PRICE_HEADER]

    updated_count = 0
    for row in range(HEADER_ROW + 1, worksheet.max_row + 1):
        sample_name = str(worksheet.cell(row=row, column=sample_column).value or "").strip()
        common_name = str(worksheet.cell(row=row, column=common_name_column).value or "").strip()
        price = prices_by_food.get((sample_name, common_name))
        if price is not None:
            worksheet.cell(row=row, column=price_column).value = price
            updated_count += 1

    workbook.save(workbook_path)
    print(f"已更新 {updated_count} 筆價格至：{workbook_path}", flush=True)


def run_pipeline(directory: Path, dry_run: bool) -> None:
    if not directory.is_dir():
        raise NotADirectoryError(f"工作資料夾不存在：{directory}")
    print(f"工作資料夾：{directory}", flush=True)
    for stage in STAGES:
        run_stage(stage, directory, dry_run)
    if not dry_run:
        update_reference_workbook(directory)
    print("\n流程完成。", flush=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="依序執行 MOA 價格與營養資料流程")
    parser.add_argument(
        "--directory",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="存放 01–08 資料處理腳本的資料夾",
    )
    parser.add_argument("--dry-run", action="store_true", help="只列出將執行的腳本")
    return parser.parse_args()


def main() -> int:
    configure_taiwan_timezone()
    print(f"執行時區：{TAIWAN_TIMEZONE}", flush=True)
    args = parse_args()
    directory = args.directory.expanduser().resolve()
    run_pipeline(directory, args.dry_run)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, NotADirectoryError, RuntimeError, subprocess.CalledProcessError, ValueError) as error:
        print(f"流程失敗：{error}", file=sys.stderr)
        raise SystemExit(1)
