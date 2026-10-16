#!/usr/bin/env python3
"""Match column-A ingredients in ytower_nutrition_comparison.xlsx to Wan Jia Fu.

This module reuses the catalog crawl, site-search fallback, matching rules,
product-spec parsing, and per-100g calculation in the adjacent wanjia_crawler.py.
It is designed for the repository's Airflow/Docker mount layout as well as local
runs. The source workbook is backed up once before the first write.

Docker/Airflow example:
  python /workspace/python/Price/wanjia_ytower_nutrition_comparison.py \
    --workbook /workspace/python/Price/ytower_nutrition_comparison.xlsx \
    --output-dir /workspace/data/processed/wanjia_ytower_comparison
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import time
from collections import defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# The script and source crawler are mounted together at /workspace/python/Price
# in Docker and are adjacent in the local repository.
import wanjia_crawler as crawler

# A's 「米飯」 and B's 「白飯」 refer to cooked rice here. Search the store's
# shelf-stable cooked-rice term as well, and prefer this verified product page
# whenever it appears in a result set.
RICE_PRODUCT_URL = f"{crawler.BASE_URL}/zh/2020109100101.html"
crawler.ALIASES.update({
    "米飯": ("常溫熟飯", "白飯"),
    "白飯": ("常溫熟飯", "米飯"),
    "常溫熟飯": ("白飯", "米飯"),
})

HEADER_ROW = 2
DEFAULT_SHEET = "食材對照"
NUTRITION_SHEET = "台灣食品成分表"
NUTRITION_PRICE_HEADER = "每100g的價格"
BASE_INGREDIENT_PREFIXES = (
    "日本", "台灣", "臺灣", "本產", "進口", "國產", "有機", "新鮮", "冷凍", "冷藏",
    "去皮", "帶皮", "去骨", "帶骨", "切片", "切丁", "切絲", "切末", "切塊",
)
NONFOOD_RESULT_RE = re.compile(r"烤麵包機|烤吐司機|吐司機|麵包機|烤箱")
OUTPUT_HEADERS = (
    "萬家福狀態",
    "萬家福商品名稱",
    "萬家福商品網址",
    "萬家福網站大類",
    "萬家福商品價格",
    "萬家福規格",
    "萬家福重量g",
    "萬家福每100g價格",
    "萬家福相似度",
    "萬家福命中搜尋詞",
    "萬家福搜尋詞",
    "萬家福未換算原因",
)
CHECKPOINT_VERSION = 1


def default_workbook() -> Path:
    configured = os.environ.get("YTOWER_COMPARISON_XLSX")
    if configured:
        return Path(configured).expanduser()
    return Path(__file__).resolve().with_name("ytower_nutrition_comparison.xlsx")


def default_nutrition_workbook() -> Path:
    configured = os.environ.get("FOOD_NUTRITION_XLSX")
    if configured:
        return Path(configured).expanduser()
    # In Docker, __file__ is /workspace/python/Price/...; locally it is
    # <repository>/python/Price/.... Both resolve to the project data folder.
    return Path(__file__).resolve().parents[2] / "data" / "reference" / "food_nutrition_2025.xlsx"


def clean(value: Any) -> str:
    return crawler.clean(value)


def search_base_name(value: str) -> str:
    """Remove common origin/state descriptors so price searches use the food itself."""
    name = crawler.base_name(value)
    previous = None
    while name != previous:
        previous = name
        for prefix in BASE_INGREDIENT_PREFIXES:
            if name.startswith(prefix) and len(name) > len(prefix):
                name = name[len(prefix):].strip()
                break
    return name or crawler.base_name(value)


def group_key(group: list[dict[str, Any]]) -> str:
    return crawler.normalize(search_base_name(group[0]["樣品名稱"]))


def checkpoint_paths(args: argparse.Namespace) -> tuple[Path, Path, Path]:
    state_path = args.checkpoint or (args.output_dir / "wanjia_ytower_comparison.checkpoint.json")
    return state_path, state_path.with_suffix(".results.jsonl"), state_path.with_suffix(".catalog.json")


def checkpoint_signature(samples: list[dict[str, Any]], args: argparse.Namespace) -> str:
    source = {
        "samples": [{key: sample.get(key) for key in
                     ("Excel列", "食品分類", "樣品名稱", "俗名", "建議搜尋詞", "營養樣品名稱空白")}
                    for sample in samples],
        "sheet": args.sheet,
        "nutrition_sheet": args.nutrition_sheet,
        "threshold": args.threshold,
        "max_pages": args.max_pages,
        "search_limit": args.search_limit,
    }
    return hashlib.sha256(json.dumps(source, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def write_json_atomic(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    os.replace(temporary, path)


def load_checkpoint(args: argparse.Namespace, samples: list[dict[str, Any]]) -> dict[str, Any]:
    state_path, journal_path, catalog_path = checkpoint_paths(args)
    signature = checkpoint_signature(samples, args)
    if args.restart:
        for path in (state_path, journal_path, catalog_path):
            if path.exists():
                path.unlink()
    state = {"version": CHECKPOINT_VERSION, "signature": signature, "searched_count": 0,
             "active": None, "network_failures": [], "resumed": False}
    if state_path.exists():
        try:
            saved = json.loads(state_path.read_text(encoding="utf-8"))
            if saved.get("version") == CHECKPOINT_VERSION and saved.get("signature") == signature:
                state.update(saved)
                state["resumed"] = True
            else:
                print("檢查點與目前食材清單或搜尋設定不同，將從頭開始。", flush=True)
                journal_path.write_text("", encoding="utf-8")
                if catalog_path.exists():
                    catalog_path.unlink()
        except (OSError, json.JSONDecodeError):
            print("檢查點無法讀取，將從頭開始。", flush=True)
            journal_path.write_text("", encoding="utf-8")
    if not state["resumed"] and journal_path.exists():
        journal_path.write_text("", encoding="utf-8")
    state["results"] = {}
    state["completed_groups"] = set()
    if state["resumed"] and journal_path.exists():
        try:
            for line in journal_path.read_text(encoding="utf-8").splitlines():
                if not line.strip():
                    continue
                entry = json.loads(line)
                key = entry["group_key"]
                for row in entry.get("rows", []):
                    state["results"][int(row["Excel列"])] = row
                if entry.get("complete"):
                    state["completed_groups"].add(key)
                else:
                    state["completed_groups"].discard(key)
        except (OSError, json.JSONDecodeError, KeyError, TypeError, ValueError):
            state["results"] = {}
            state["completed_groups"] = set()
            state["active"] = None
            state["resumed"] = False
            journal_path.write_text("", encoding="utf-8")
            print("結果記錄不完整，將從頭開始。", flush=True)
    state["state_path"], state["journal_path"], state["catalog_path"] = state_path, journal_path, catalog_path
    state["signature"] = signature
    save_checkpoint_state(state)
    if state["resumed"]:
        print(f"續跑檢查點：已完成 {len(state['completed_groups'])} 組；搜尋請求累計 {state['searched_count']} 次。", flush=True)
    return state


def save_checkpoint_state(state: dict[str, Any]) -> None:
    persisted = {key: value for key, value in state.items()
                 if key not in {"results", "completed_groups", "state_path", "journal_path", "catalog_path"}}
    write_json_atomic(state["state_path"], persisted)


def save_group_checkpoint(state: dict[str, Any], key: str,
                          rows: list[dict[str, Any]], complete: bool) -> None:
    with state["journal_path"].open("a", encoding="utf-8") as stream:
        stream.write(json.dumps({"group_key": key, "rows": rows, "complete": complete}, ensure_ascii=False) + "\n")
        stream.flush()
        os.fsync(stream.fileno())
    for row in rows:
        state["results"][int(row["Excel列"])] = row
    if complete:
        state["completed_groups"].add(key)
    else:
        state["completed_groups"].discard(key)
    save_checkpoint_state(state)


def read_ingredients(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    """Read A-column ingredient rows and capture B/C/E as extra search context."""
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError(f"工作簿找不到工作表「{sheet_name}」")
        sheet = book[sheet_name]
        ingredients: list[dict[str, Any]] = []
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True),
            start=HEADER_ROW + 1,
        ):
            row = list(values)
            name = clean(row[0] if row else "")
            if not name:
                continue
            # Category divider rows have only a label in A; actual ingredient
            # rows carry their nutrition category in E, even when B-D are blank.
            has_other_content = any(
                clean(row[index])
                for index in (1, 2, 3, 4)
                if index < len(row) and row[index] is not None
            )
            if not has_other_content:
                continue
            ingredients.append({
                "Excel列": excel_row,
                "來源列號": str(excel_row),
                "食品分類": clean(row[4] if len(row) > 4 else ""),
                "樣品名稱": name,
                "基礎食材": search_base_name(name),
                "營養樣品名稱空白": not clean(row[1] if len(row) > 1 else ""),
                # Existing sample name and alias columns are helpful additional
                # search terms; column A remains the term being priced.
                "俗名": clean(row[2] if len(row) > 2 else ""),
                "建議搜尋詞": clean(row[1] if len(row) > 1 else ""),
            })
        return ingredients
    finally:
        book.close()


def ingredient_groups(samples: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for sample in samples:
        key = crawler.normalize(search_base_name(sample["樣品名稱"]))
        if key:
            grouped.setdefault(key, []).append(sample)
    return list(grouped.values())


def representative(group: list[dict[str, Any]]) -> dict[str, Any]:
    result = group[0].copy()
    aliases = []
    searches = []
    for sample in group:
        for value, destination in ((sample.get("俗名"), aliases), (sample.get("建議搜尋詞"), searches)):
            value = clean(value)
            if value and value not in destination:
                destination.append(value)
    result["俗名"] = "、".join(aliases)
    result["建議搜尋詞"] = "、".join(searches)
    result["樣品名稱"] = search_base_name(result["樣品名稱"])
    return result


def prioritize_rice_product(products: list[dict[str, Any]], sample: dict[str, Any]) -> list[dict[str, Any]]:
    """Put the requested cooked-rice product first for rice ingredient groups."""
    values = " ".join(clean(sample.get(field)) for field in ("樣品名稱", "俗名", "建議搜尋詞"))
    is_rice = any(crawler.normalize(term) in crawler.normalize(values)
                  for term in ("米飯", "白飯", "常溫熟飯"))
    if not is_rice:
        return products
    preferred = [p for p in products if clean(p.get("商品網址", p.get("網址", ""))).rstrip("/") == RICE_PRODUCT_URL.rstrip("/")]
    if not preferred:
        return products
    preferred_ids = {id(p) for p in preferred}
    return preferred + [p for p in products if id(p) not in preferred_ids]


def fan_out(match: dict[str, Any], group: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for sample in group:
        row = dict(match)
        row.update({
            key: sample.get(key, "")
            for key in ("來源列號", "食品分類", "樣品名稱", "俗名")
        })
        row["Excel列"] = sample["Excel列"]
        row["補查基礎食材"] = search_base_name(sample["樣品名稱"])
        row["站內搜尋詞"] = row.get("站內搜尋詞", [])
        rows.append(row)
    return rows


def no_match_row(sample: dict[str, Any], attempted: list[str], failed_terms: list[str],
                 network_failures: list[dict[str, Any]]) -> dict[str, Any]:
    catalog_failed = any(error.get("階段") == "型錄" for error in network_failures)
    if attempted and len(failed_terms) == len(attempted):
        reason = "所有已嘗試的站內搜尋都因網路/DNS解析失敗，未能確認是否有商品；請恢復網路後重跑"
    elif failed_terms:
        reason = "型錄與成功完成的站內搜尋未找到符合候選；部分搜尋遇到網路/DNS解析失敗，建議重跑確認"
    elif catalog_failed:
        reason = "未找到符合候選；本次有型錄頁因網路/DNS失敗未讀取，請恢復網路後重跑確認"
    else:
        reason = "分類型錄及已嘗試的站內搜尋詞均未找到符合候選"
    return {
        **{key: sample.get(key, "") for key in ("來源列號", "食品分類", "樣品名稱", "俗名")},
        "狀態": "未找到",
        "補查基礎食材": search_base_name(sample["樣品名稱"]),
        "商品名稱": "", "商品網址": "", "網站大類": "", "商品價格": None,
        "規格": "", "重量g": None, "重量估算說明": "", "每100g的價格": None,
        "密度假設": "", "相似度": None, "命中搜尋詞": "",
        "站內搜尋詞": attempted, "未換算原因": reason,
        "Excel列": sample["Excel列"],
    }


def external_search_results(session, query: str, timeout: float) -> list[dict[str, str]]:
    """Read ordinary web-search titles/snippets for conservative measure/price fallbacks."""
    url = "https://html.duckduckgo.com/html/?" + urlencode({"q": query})
    response = session.get(url, timeout=timeout, headers={"Accept-Language": "zh-TW,zh;q=0.9"})
    response.raise_for_status()
    soup = crawler.BeautifulSoup(response.text, "html.parser")
    results = []
    for result in soup.select(".result"):
        title_node = result.select_one(".result__title")
        snippet_node = result.select_one(".result__snippet")
        title = clean(title_node.get_text(" ", strip=True)) if title_node else ""
        snippet = clean(snippet_node.get_text(" ", strip=True)) if snippet_node else ""
        if title or snippet:
            results.append({"標題": title, "摘要": snippet})
    return results


def web_measure_fallback(session, ingredient: str, product_name: str, timeout: float,
                         delay: float, failures: list[dict[str, Any]]) -> dict[str, Any] | None:
    queries = (
        f'"{product_name}" 內容量',
        f'"{ingredient}" "{product_name}" 重量 體積',
    )
    for query in queries:
        try:
            results = external_search_results(session, query, timeout)
        except Exception as exc:
            failures.append({"階段": "網路重量補查", "食材": ingredient, "搜尋詞": query, "錯誤": str(exc)})
            time.sleep(max(0.0, delay))
            continue
        for item in results:
            text = f"{item['標題']} {item['摘要']}"
            measure = crawler.extract_measure(text)
            key = crawler.normalize(ingredient)
            if key and key in crawler.normalize(text) and measure.get("重量g"):
                return {**measure, "來源": "網路搜尋摘要", "搜尋詞": query, "摘要": text}
        time.sleep(max(0.0, delay))
    return None


def web_price_fallback(session, ingredient: str, timeout: float, delay: float,
                       failures: list[dict[str, Any]]) -> dict[str, Any] | None:
    queries = (f'"{ingredient}" 價格 內容量', f'"{ingredient}" 重量 體積 價格')
    for query in queries:
        try:
            results = external_search_results(session, query, timeout)
        except Exception as exc:
            failures.append({"階段": "網路價格補查", "食材": ingredient, "搜尋詞": query, "錯誤": str(exc)})
            time.sleep(max(0.0, delay))
            continue
        for item in results:
            text = f"{item['標題']} {item['摘要']}"
            if not ingredient or crawler.normalize(ingredient) not in crawler.normalize(text):
                continue
            measure = crawler.extract_measure(text)
            price_match = re.search(r"(?:NT\s*\$?|TWD\s*|(?:售價|價格|特價|市價)\s*[:：]?\s*|\$\s*)([\d,]+(?:\.\d+)?)\s*(?:元|NTD)?", text, re.I)
            if not price_match or not measure.get("重量g"):
                continue
            price = float(price_match.group(1).replace(",", ""))
            if price <= 0:
                continue
            per_100g = round(price / measure["重量g"] * 100, 4)
            return {
                "商品名稱": item["標題"], "商品網址": "", "商品價格": price,
                "規格": measure.get("規格", ""), "重量g": measure["重量g"],
                "每100g的價格": per_100g, "補查來源": "網路搜尋摘要",
                "補查搜尋詞": query, "補查摘要": text,
            }
        time.sleep(max(0.0, delay))
    return None


def complete_missing_weight(match: dict[str, Any], sample: dict[str, Any], session,
                            args: argparse.Namespace, network_failures: list[dict[str, Any]]) -> None:
    if match.get("重量g") or NONFOOD_RESULT_RE.search(clean(match.get("商品名稱"))):
        return
    measure = web_measure_fallback(
        session, search_base_name(sample["樣品名稱"]), clean(match.get("商品名稱")),
        args.timeout, args.delay, network_failures,
    )
    if not measure:
        return
    match["重量g"] = measure["重量g"]
    match["規格"] = measure.get("規格", "")
    match["每100g的價格"] = crawler.price_per_100g(
        {"商品價格": match.get("商品價格")}, measure,
    )[0]
    match["狀態"] = "成功" if match["每100g的價格"] is not None else "找到商品但未換算"
    match["未換算原因"] = "" if match["每100g的價格"] is not None else "網路補查重量後仍缺少萬家福商品價格"
    match["重量估算說明"] = "網路搜尋摘要補查；請核對來源"
    match["密度假設"] = measure.get("密度假設", "")
    match["重量補查來源"] = f"{measure['來源']}：{measure['搜尋詞']}"


def write_back_nutrition_workbook(nutrition_path: Path, samples: list[dict[str, Any]],
                                  results: list[dict[str, Any]], sheet_name: str) -> tuple[int, int, Path]:
    if not nutrition_path.is_file():
        raise FileNotFoundError(f"找不到原始營養資料工作簿：{nutrition_path}")
    backup_path = nutrition_path.with_name(nutrition_path.stem + ".before_wanjia_price.xlsx")
    if not backup_path.exists():
        shutil.copy2(nutrition_path, backup_path)
    book = load_workbook(nutrition_path)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError(f"原始營養資料找不到工作表「{sheet_name}」")
        sheet = book[sheet_name]
        headers = {clean(sheet.cell(HEADER_ROW, col).value): col for col in range(1, sheet.max_column + 1)}
        name_col = headers.get("樣品名稱")
        category_col = headers.get("食品分類")
        alias_col = headers.get("俗名")
        if not name_col:
            raise ValueError("原始營養資料缺少「樣品名稱」欄")
        price_col = headers.get(NUTRITION_PRICE_HEADER)
        if not price_col:
            price_col = sheet.max_column + 1
            sheet.cell(HEADER_ROW, price_col).value = NUTRITION_PRICE_HEADER
        existing: dict[str, list[int]] = defaultdict(list)
        existing_exact: set[str] = set()
        for row_number in range(HEADER_ROW + 1, sheet.max_row + 1):
            name = clean(sheet.cell(row_number, name_col).value)
            if name:
                existing[crawler.normalize(search_base_name(name))].append(row_number)
                existing_exact.add(crawler.normalize(name))
        result_by_base: dict[str, dict[str, Any]] = {}
        for result in results:
            base = search_base_name(clean(result.get("補查基礎食材") or result.get("樣品名稱")))
            if base and result.get("每100g的價格") is not None:
                result_by_base[crawler.normalize(base)] = result
        updated = 0
        for key, rows in existing.items():
            result = result_by_base.get(key)
            if not result:
                continue
            for row_number in rows:
                sheet.cell(row_number, price_col).value = result["每100g的價格"]
                sheet.cell(row_number, price_col).number_format = "#,##0.0000"
                updated += 1
        added = 0
        seen = set(existing_exact)
        for sample in samples:
            if not sample.get("營養樣品名稱空白"):
                continue
            name = clean(sample.get("樣品名稱"))
            key = crawler.normalize(name)
            if not key or key in seen:
                continue
            seen.add(key)
            row_number = sheet.max_row + 1
            if category_col:
                sheet.cell(row_number, category_col).value = sample.get("食品分類") or "待補營養資料"
            sheet.cell(row_number, name_col).value = name
            if alias_col and sample.get("俗名"):
                sheet.cell(row_number, alias_col).value = sample["俗名"]
            result = result_by_base.get(crawler.normalize(search_base_name(name)))
            if result:
                sheet.cell(row_number, price_col).value = result["每100g的價格"]
                sheet.cell(row_number, price_col).number_format = "#,##0.0000"
            added += 1
        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(prefix=f"{nutrition_path.stem}_", suffix=".xlsx",
                                             dir=nutrition_path.parent, delete=False) as temp:
                temporary_path = Path(temp.name)
            book.save(temporary_path)
            os.replace(temporary_path, nutrition_path)
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()
        return updated, added, backup_path
    finally:
        book.close()


def crawl_and_match(samples: list[dict[str, Any]], args: argparse.Namespace) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    state = load_checkpoint(args, samples)
    session = crawler.create_session()
    network_failures: list[dict[str, Any]] = list(state.get("network_failures", []))
    products = None
    if state["resumed"] and state["catalog_path"].exists():
        try:
            cached = json.loads(state["catalog_path"].read_text(encoding="utf-8"))
            if cached.get("signature") == state["signature"]:
                products = cached.get("products", [])
                print(f"使用檢查點型錄：{len(products)} 筆商品。", flush=True)
        except (OSError, json.JSONDecodeError):
            products = None
    if products is None:
        # Catalog data is cached only when every category page was read; an
        # incomplete crawl is retried after a later restart with network access.
        before_catalog_failures = len(network_failures)
        products = crawler.crawl_catalog(
            session, args.max_pages, args.timeout, args.delay, network_failures,
        )
        if len(network_failures) == before_catalog_failures:
            write_json_atomic(state["catalog_path"], {"signature": state["signature"], "products": products})
    products = [product for product in products if not NONFOOD_RESULT_RE.search(clean(product.get("商品名稱")))]
    spec_cache: dict[str, tuple[str, dict[str, Any]]] = {}
    groups = ingredient_groups(samples)
    results: dict[int, dict[str, Any]] = dict(state["results"])
    unmatched: list[list[dict[str, Any]]] = []

    # Preserve the source crawler's first pass: match against crawled category
    # listings before falling back to individual site searches.
    for group in groups:
        key = group_key(group)
        if key in state["completed_groups"]:
            print(f"跳過已完成食材群：{search_base_name(group[0]['樣品名稱'])}", flush=True)
            continue
        search_sample = representative(group)
        match = crawler.match_from_pool(
            search_sample, prioritize_rice_product(products, search_sample), args.threshold, session, args.timeout, spec_cache,
        )
        if match:
            complete_missing_weight(match, search_sample, session, args, network_failures)
            match["站內搜尋詞"] = []
            rows = fan_out(match, group)
            for row in rows:
                results[row["Excel列"]] = row
            state["active"] = None
            state["network_failures"] = network_failures
            save_group_checkpoint(state, key, rows, True)
        else:
            unmatched.append(group)

    print(f"型錄未匹配食材群：{len(unmatched)} 組；開始逐組站內搜尋。", flush=True)
    searched = int(state.get("searched_count", 0))
    for group in unmatched:
        key = group_key(group)
        if key in state["completed_groups"]:
            continue
        search_group = [{**sample, "樣品名稱": search_base_name(sample["樣品名稱"])} for sample in group]
        active_state = state.get("active") or {}
        active = active_state if active_state.get("group_key") == key else None

        def save_search_progress(next_index: int, pending_indices: list[int], attempted_terms: list[str],
                                 failed_terms: list[str], request_count: int) -> None:
            state["active"] = {"group_key": key, "next_index": next_index,
                               "pending_indices": pending_indices, "attempted": attempted_terms,
                               "failed_terms": failed_terms}
            state["searched_count"] = request_count
            state["network_failures"] = network_failures
            save_checkpoint_state(state)

        found, searched, attempted, failed_terms = site_search_food_group(
            session, search_group, args, searched, spec_cache, network_failures,
            active, save_search_progress,
        )
        if found:
            complete_missing_weight(found[0], search_group[0], session, args, network_failures)
            for row in found:
                original = next(sample for sample in group if sample["Excel列"] == row["Excel列"])
                row.update({key: original.get(key, "") for key in ("食品分類", "樣品名稱", "俗名")})
                row["補查基礎食材"] = search_base_name(original["樣品名稱"])
                for key in ("重量g", "規格", "每100g的價格", "狀態", "未換算原因", "重量估算說明", "重量補查來源"):
                    if key in found[0]:
                        row[key] = found[0][key]
                results[row["Excel列"]] = row
            state["active"] = None
            state["network_failures"] = network_failures
            save_group_checkpoint(state, group_key(group), found, True)
        else:
            adjective_group = any(search_base_name(sample["樣品名稱"]) != crawler.base_name(sample["樣品名稱"]) for sample in group)
            web_match = None
            if adjective_group:
                web_match = web_price_fallback(
                    session, search_base_name(group[0]["樣品名稱"]), args.timeout, args.delay, network_failures,
                )
            for sample in group:
                row = no_match_row(sample, attempted, failed_terms, network_failures)
                if web_match:
                    row.update(web_match)
                    row["狀態"] = "網路搜尋估價"
                    row["補查基礎食材"] = search_base_name(sample["樣品名稱"])
                    row["未換算原因"] = f"萬家福未找到；依基礎食材網路搜尋估價，請核對：{web_match['補查搜尋詞']}"
                results[row["Excel列"]] = row
            active_after_search = state.get("active") or {}
            term_count = len(crawler.sample_search_terms(representative(search_group)))
            retryable = bool(failed_terms) or (
                active_after_search.get("group_key") == key
                and (active_after_search.get("pending_indices") or active_after_search.get("next_index", 0) < term_count)
            )
            if not retryable:
                state["active"] = None
            state["network_failures"] = network_failures
            state["searched_count"] = searched
            save_group_checkpoint(state, key, [results[sample["Excel列"]] for sample in group], not retryable)

    ordered = [results[sample["Excel列"]] for sample in samples]
    summary = {
        "型錄商品數": len(products),
        "型錄與站內搜尋連線失敗筆數": len(network_failures),
        "連線失敗紀錄": network_failures,
        "站內搜尋請求數": searched,
        "已由檢查點續跑": state["resumed"],
        "檢查點已完成食材群": len(state["completed_groups"]),
        "食材列數": len(samples),
        "不同食材群數": len(groups),
        "成功換算每100g筆數": sum(row.get("每100g的價格") is not None for row in ordered),
        "找到商品但未換算筆數": sum(row.get("狀態") == "找到商品但未換算" for row in ordered),
        "未找到筆數": sum(row.get("狀態") == "未找到" for row in ordered),
        "網路補查重量筆數": sum(bool(row.get("重量補查來源")) for row in ordered),
        "網路搜尋估價筆數": sum(row.get("狀態") == "網路搜尋估價" for row in ordered),
    }
    return ordered, summary


def site_search_food_group(session, group: list[dict[str, Any]], args: argparse.Namespace,
                           searched_count: int, spec_cache: dict[str, Any],
                           network_failures: list[dict[str, Any]], active: dict[str, Any] | None,
                           progress_callback) -> tuple[list[dict[str, Any]] | None, int, list[str], list[str]]:
    terms, seen = [], set()
    for sample in group:
        for term in crawler.sample_search_terms(sample):
            key = crawler.normalize(term)
            if key and key not in seen:
                terms.append(term)
                seen.add(key)
    attempted = list((active or {}).get("attempted", []))
    failed_terms = list((active or {}).get("failed_terms", []))
    representative_sample = representative(group)
    next_index = int((active or {}).get("next_index", 0))
    pending_indices = sorted(set(int(index) for index in (active or {}).get("pending_indices", [])))

    def run_term(index: int) -> dict[str, Any] | None:
        nonlocal searched_count
        term = terms[index]
        if args.search_limit and searched_count >= args.search_limit:
            return None
        if term not in attempted:
            attempted.append(term)
        url = f"{crawler.BASE_URL}/zh/search/?{urlencode({'q': term})}"
        print(f"站內搜尋 {searched_count + 1}{'/' + str(args.search_limit) if args.search_limit else ''}：{representative_sample['樣品名稱']}｜{term}", flush=True)
        searched_count += 1
        try:
            html, final_url = crawler.get_html(session, url, args.timeout)
            products = [product for product in crawler.extract_products(html, final_url)
                        if not NONFOOD_RESULT_RE.search(clean(product.get("商品名稱")))]
            match = crawler.match_from_pool(representative_sample,
                                            prioritize_rice_product(products, representative_sample), args.threshold,
                                            session, args.timeout, spec_cache)
        except Exception as exc:
            if term not in failed_terms:
                failed_terms.append(term)
            network_failures.append({"階段": "站內搜尋", "食材": representative_sample["樣品名稱"],
                                     "搜尋詞": term, "網址": url, "錯誤": str(exc)})
            if index not in pending_indices:
                pending_indices.append(index)
                pending_indices.sort()
            progress_callback(next_index, pending_indices, attempted, failed_terms, searched_count)
            time.sleep(max(0.0, args.delay))
            return None
        if term in failed_terms:
            failed_terms.remove(term)
        if index in pending_indices:
            pending_indices.remove(index)
        return match

    for index in pending_indices.copy():
        if args.search_limit and searched_count >= args.search_limit:
            progress_callback(next_index, pending_indices, attempted, failed_terms, searched_count)
            return None, searched_count, attempted, failed_terms
        match = run_term(index)
        if match:
            rows = []
            for sample in group:
                row = {**match, **{key: sample.get(key, "") for key in ("來源列號", "食品分類", "樣品名稱", "俗名")}}
                row["Excel列"] = sample["Excel列"]
                row["補查基礎食材"] = search_base_name(sample["樣品名稱"])
                row["站內搜尋詞"] = attempted.copy()
                rows.append(row)
            return rows, searched_count, attempted, failed_terms
        progress_callback(next_index, pending_indices, attempted, failed_terms, searched_count)
        time.sleep(max(0.0, args.delay))
    for index in range(next_index, len(terms)):
        if args.search_limit and searched_count >= args.search_limit:
            progress_callback(index, pending_indices, attempted, failed_terms, searched_count)
            return None, searched_count, attempted, failed_terms
        match = run_term(index)
        if match:
            rows = []
            for sample in group:
                row = {**match, **{key: sample.get(key, "") for key in ("來源列號", "食品分類", "樣品名稱", "俗名")}}
                row["Excel列"] = sample["Excel列"]
                row["補查基礎食材"] = search_base_name(sample["樣品名稱"])
                row["站內搜尋詞"] = attempted.copy()
                rows.append(row)
            return rows, searched_count, attempted, failed_terms
        next_index = index + 1
        progress_callback(next_index, pending_indices, attempted, failed_terms, searched_count)
        time.sleep(max(0.0, args.delay))
    return None, searched_count, attempted, failed_terms


def output_column_start(sheet) -> int:
    headers = [clean(sheet.cell(HEADER_ROW, col).value) for col in range(1, max(sheet.max_column, 5) + 1)]
    expected = list(OUTPUT_HEADERS)
    if headers[5:5 + len(expected)] == expected:
        return 6
    if not any(headers[5:]):
        return 6
    # Idempotently update an already appended block even if unrelated columns
    # were added later, otherwise never overwrite user-owned columns.
    for start in range(6, len(headers) - len(expected) + 2):
        if headers[start - 1:start - 1 + len(expected)] == expected:
            return start
    return max(6, sheet.max_column + 1)


def write_results(workbook_path: Path, sheet_name: str,
                  results: list[dict[str, Any]], samples: list[dict[str, Any]]) -> tuple[int, Path]:
    backup_path = workbook_path.with_name(workbook_path.stem + ".before_wanjia_comparison.xlsx")
    if not backup_path.exists():
        shutil.copy2(workbook_path, backup_path)
        print(f"已建立首次更新備份：{backup_path}", flush=True)
    book = load_workbook(workbook_path)
    if sheet_name not in book.sheetnames:
        book.close()
        raise ValueError(f"工作簿找不到工作表「{sheet_name}」")
    sheet = book[sheet_name]
    start_col = output_column_start(sheet)
    header_template = sheet.cell(HEADER_ROW, 5)
    for offset, title in enumerate(OUTPUT_HEADERS):
        cell = sheet.cell(HEADER_ROW, start_col + offset)
        if header_template.has_style:
            cell._style = header_template._style
        cell.value = title
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    results_by_row = {result["Excel列"]: result for result in results}
    sample_rows = {sample["Excel列"] for sample in samples}
    for row_number in sample_rows:
        result = results_by_row[row_number]
        if result.get("每100g的價格") is not None:
            sheet.cell(row_number, 4).value = result["每100g的價格"]
            sheet.cell(row_number, 4).number_format = "#,##0.0000"
        values = (
            result.get("狀態", ""), result.get("商品名稱", ""), result.get("商品網址", ""),
            result.get("網站大類", ""), result.get("商品價格"), result.get("規格", ""),
            result.get("重量g"), result.get("每100g的價格"), result.get("相似度"),
            result.get("命中搜尋詞", ""), "、".join(result.get("站內搜尋詞", [])),
            result.get("未換算原因", ""),
        )
        for offset, value in enumerate(values):
            cell = sheet.cell(row_number, start_col + offset)
            template = sheet.cell(row_number, 5)
            if template.has_style:
                cell._style = template._style
            cell.value = value
            cell.alignment = Alignment(vertical="center", wrap_text=offset in {1, 2, 5, 9, 10, 11})
        sheet.cell(row_number, start_col + 4).number_format = "#,##0.00"
        sheet.cell(row_number, start_col + 6).number_format = "#,##0.00"
        sheet.cell(row_number, start_col + 7).number_format = "0.0000"
        sheet.cell(row_number, start_col + 8).number_format = "0.00%"

    widths = (18, 32, 48, 16, 16, 24, 14, 18, 14, 20, 48, 52)
    for offset, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(start_col + offset)].width = width

    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f"{workbook_path.stem}_", suffix=".xlsx",
            dir=workbook_path.parent, delete=False,
        ) as temp:
            temporary_path = Path(temp.name)
        book.save(temporary_path)
        book.close()
        os.replace(temporary_path, workbook_path)
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink()
    print(f"已將萬家福比對欄位寫入 {workbook_path}（起始欄：{get_column_letter(start_col)}）", flush=True)
    return start_col, backup_path


def write_reports(output_dir: Path, workbook_path: Path, results: list[dict[str, Any]],
                  summary: dict[str, Any], args: argparse.Namespace) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "wanjia_ytower_comparison.json"
    csv_path = output_dir / "wanjia_ytower_comparison.csv"
    report = {
        "工作簿": str(workbook_path), "工作表": args.sheet,
        "比對門檻": args.threshold,
        "分類網址": crawler.CATEGORY_URLS,
        **summary,
        "結果": [{key: value for key, value in row.items() if key != "Excel列"} for row in results],
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    fields = [
        "來源列號", "食品分類", "樣品名稱", "補查基礎食材", "俗名", "狀態", "商品名稱", "商品網址",
        "網站大類", "商品價格", "規格", "重量g", "每100g的價格", "相似度",
        "命中搜尋詞", "站內搜尋詞", "重量補查來源", "補查來源", "補查搜尋詞", "未換算原因",
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in results:
            writer.writerow({**row, "站內搜尋詞": "、".join(row.get("站內搜尋詞", []))})
    return report_path, csv_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="比對食材營養對照表A欄與萬家福商品，並將結果寫入F欄之後",
    )
    parser.add_argument("--workbook", type=Path, default=default_workbook(),
                        help="ytower_nutrition_comparison.xlsx 路徑")
    parser.add_argument("--nutrition-workbook", type=Path, default=default_nutrition_workbook(),
                        help="food_nutrition_2025.xlsx 原始營養資料路徑")
    parser.add_argument("--nutrition-sheet", default=os.environ.get("FOOD_NUTRITION_SHEET", NUTRITION_SHEET))
    parser.add_argument("--sheet", default=os.environ.get("YTOWER_COMPARISON_SHEET", DEFAULT_SHEET))
    parser.add_argument(
        "--output-dir", type=Path,
        default=Path(os.environ.get(
            "WANJIA_COMPARISON_OUTPUT_DIR",
            str(Path(__file__).resolve().parents[2] / "data" / "processed" / "wanjia_ytower_comparison"),
        )),
    )
    parser.add_argument("--threshold", type=float, default=float(os.environ.get("WANJIA_THRESHOLD", "0.60")))
    parser.add_argument("--delay", type=float, default=float(os.environ.get("WANJIA_DELAY", "1.0")))
    parser.add_argument("--timeout", type=float, default=float(os.environ.get("WANJIA_TIMEOUT", "25")))
    parser.add_argument("--max-pages", type=int, default=int(os.environ.get("WANJIA_MAX_PAGES", "0")))
    parser.add_argument("--search-limit", type=int, default=int(os.environ.get("WANJIA_SEARCH_LIMIT", "0")))
    parser.add_argument("--checkpoint", type=Path, help="續跑檢查點 JSON 路徑（預設存於 --output-dir）")
    parser.add_argument("--restart", action="store_true", help="忽略並清除上一輪檢查點，從頭搜尋")
    parser.add_argument("--no-update-workbook", action="store_true", help="只輸出 JSON/CSV 報告，不修改活頁簿")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not 0 <= args.threshold <= 1:
        raise SystemExit("--threshold 必須介於 0 到 1")
    if args.timeout <= 0 or args.delay < 0 or args.max_pages < 0 or args.search_limit < 0:
        raise SystemExit("timeout 必須大於0，delay/max-pages/search-limit 不可小於0")
    if not args.workbook.is_file():
        raise SystemExit(f"找不到食材對照工作簿：{args.workbook}")
    if not args.no_update_workbook and not args.nutrition_workbook.is_file():
        raise SystemExit(f"找不到原始營養資料工作簿：{args.nutrition_workbook}")
    samples = read_ingredients(args.workbook, args.sheet)
    if not samples:
        raise SystemExit("A欄沒有找到可比對的食材資料")
    print(f"讀取A欄食材：{len(samples)} 列、{len(ingredient_groups(samples))} 組；工作簿：{args.workbook}", flush=True)

    results, summary = crawl_and_match(samples, args)
    if not args.no_update_workbook:
        nutrition_updated, nutrition_added, nutrition_backup = write_back_nutrition_workbook(
            args.nutrition_workbook, samples, results, args.nutrition_sheet,
        )
        summary["原始營養資料更新價格列數"] = nutrition_updated
        summary["原始營養資料新增食材列數"] = nutrition_added
        summary["原始營養資料首次更新備份"] = str(nutrition_backup)
        start_col, backup_path = write_results(args.workbook, args.sheet, results, samples)
        summary["回填欄位起始欄"] = get_column_letter(start_col)
        summary["首次更新備份"] = str(backup_path)
    else:
        summary["回填欄位起始欄"] = "未寫入"

    report_path, csv_path = write_reports(args.output_dir, args.workbook, results, summary, args)
    print(f"JSON報告：{report_path}", flush=True)
    print(f"CSV報告：{csv_path}", flush=True)
    print(
        f"每100g成功換算：{summary['成功換算每100g筆數']}；"
        f"商品未換算：{summary['找到商品但未換算筆數']}；未找到：{summary['未找到筆數']}；"
        f"營養資料更新價格：{summary.get('原始營養資料更新價格列數', 0)}；"
        f"新增食材：{summary.get('原始營養資料新增食材列數', 0)}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
